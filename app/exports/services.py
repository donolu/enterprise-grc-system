import io
import csv
import json
import zipfile
from datetime import date, datetime, time
from decimal import Decimal
from uuid import UUID
from django.template.loader import render_to_string
from django.utils import timezone
from django.db.models import Count, Q
from django.core.files.base import ContentFile
from django.apps import apps
from django.db import models
from openpyxl import Workbook
from weasyprint import HTML, CSS
from weasyprint.text.fonts import FontConfiguration

from core.models import Document
from catalogs.models import ControlAssessment, Framework, AssessmentEvidence
from risk.analytics import RiskAnalyticsService, RiskReportGenerator
from risk.models import Risk, RiskAction
from .models import AssessmentReport, TenantDataExport


EXCLUDED_FIELD_NAMES = {
    'password',
    'last_login',
    'is_superuser',
    'is_staff',
    'user_permissions',
    'groups',
    'token',
    'token_hash',
    'secret',
    'api_key',
    'stripe_customer_id',
    'stripe_subscription_id',
    'stripe_event_id',
    'stripe_price_id',
}

EXPORT_COVERAGE = [
    {
        'module': 'identity',
        'label': 'Users, documents and audit activity',
        'formats': ['xlsx', 'csv_zip'],
        'sheets': [
            {'model': 'core.User', 'worksheet': 'Users'},
            {'model': 'core.Document', 'worksheet': 'Documents'},
            {'model': 'core.DocumentAccess', 'worksheet': 'Document access'},
            {'model': 'core.AuditEvent', 'worksheet': 'Audit events'},
        ],
    },
    {
        'module': 'frameworks',
        'label': 'Frameworks, controls, assessments, evidence and templates',
        'formats': ['xlsx', 'csv_zip'],
        'sheets': [
            {'model': 'catalogs.Framework', 'worksheet': 'Frameworks'},
            {'model': 'catalogs.Clause', 'worksheet': 'Clauses'},
            {'model': 'catalogs.Control', 'worksheet': 'Controls'},
            {'model': 'catalogs.ControlEvidence', 'worksheet': 'Control evidence'},
            {'model': 'catalogs.ControlAssessment', 'worksheet': 'Assessments'},
            {'model': 'catalogs.AssessmentEvidence', 'worksheet': 'Assessment evidence'},
            {'model': 'catalogs.TemplateDocument', 'worksheet': 'Template documents'},
            {'model': 'catalogs.AssessmentReminderLog', 'worksheet': 'Assessment reminders'},
        ],
    },
    {
        'module': 'risk',
        'label': 'Risk register, actions, evidence and reminders',
        'formats': ['xlsx', 'csv_zip'],
        'sheets': [
            {'model': 'risk.RiskCategory', 'worksheet': 'Risk categories'},
            {'model': 'risk.RiskMatrix', 'worksheet': 'Risk matrices'},
            {'model': 'risk.Risk', 'worksheet': 'Risks'},
            {'model': 'risk.RiskNote', 'worksheet': 'Risk notes'},
            {'model': 'risk.RiskAction', 'worksheet': 'Risk actions'},
            {'model': 'risk.RiskActionNote', 'worksheet': 'Risk action notes'},
            {'model': 'risk.RiskActionEvidence', 'worksheet': 'Risk action evidence'},
            {'model': 'risk.RiskActionReminderLog', 'worksheet': 'Risk action reminders'},
        ],
    },
    {
        'module': 'vendors',
        'label': 'Vendor management, contacts, services, notes and tasks',
        'formats': ['xlsx', 'csv_zip'],
        'sheets': [
            {'model': 'vendors.RegionalConfig', 'worksheet': 'Vendor regions'},
            {'model': 'vendors.VendorCategory', 'worksheet': 'Vendor categories'},
            {'model': 'vendors.Vendor', 'worksheet': 'Vendors'},
            {'model': 'vendors.VendorContact', 'worksheet': 'Vendor contacts'},
            {'model': 'vendors.VendorService', 'worksheet': 'Vendor services'},
            {'model': 'vendors.VendorNote', 'worksheet': 'Vendor notes'},
            {'model': 'vendors.VendorTask', 'worksheet': 'Vendor tasks'},
        ],
    },
    {
        'module': 'policies',
        'label': 'Policy repository, versions, acknowledgements and distributions',
        'formats': ['xlsx', 'csv_zip'],
        'sheets': [
            {'model': 'policies.PolicyCategory', 'worksheet': 'Policy categories'},
            {'model': 'policies.Policy', 'worksheet': 'Policies'},
            {'model': 'policies.PolicyVersion', 'worksheet': 'Policy versions'},
            {'model': 'policies.PolicyVersionAuditLog', 'worksheet': 'Policy audits'},
            {'model': 'policies.PolicyAcknowledgment', 'worksheet': 'Policy acknowledgements'},
            {'model': 'policies.PolicyDistribution', 'worksheet': 'Policy distributions'},
        ],
    },
    {
        'module': 'training',
        'label': 'Security awareness training, campaigns and completion tracking',
        'formats': ['xlsx', 'csv_zip'],
        'sheets': [
            {'model': 'training.TrainingCategory', 'worksheet': 'Training categories'},
            {'model': 'training.TrainingVideo', 'worksheet': 'Training videos'},
            {'model': 'training.SecurityAwarenessCampaign', 'worksheet': 'Awareness campaigns'},
            {'model': 'training.CampaignDelivery', 'worksheet': 'Campaign deliveries'},
            {'model': 'training.VideoView', 'worksheet': 'Video views'},
        ],
    },
    {
        'module': 'knowledge',
        'label': 'Knowledge base categories, articles and article revisions',
        'formats': ['xlsx', 'csv_zip'],
        'sheets': [
            {'model': 'knowledge.KnowledgeCategory', 'worksheet': 'Knowledge categories'},
            {'model': 'knowledge.KnowledgeArticle', 'worksheet': 'Knowledge articles'},
            {'model': 'knowledge.KnowledgeArticleRevision', 'worksheet': 'Knowledge revisions'},
        ],
    },
    {
        'module': 'assets',
        'label': 'Information assets and review reminders',
        'formats': ['xlsx', 'csv_zip'],
        'sheets': [
            {'model': 'assets.Asset', 'worksheet': 'Assets'},
            {'model': 'assets.AssetReviewReminderLog', 'worksheet': 'Asset reminders'},
        ],
    },
    {
        'module': 'calendar',
        'label': 'Calendar events, reminder logs and audit activity',
        'formats': ['xlsx', 'csv_zip'],
        'sheets': [
            {'model': 'calendarhub.CalendarEvent', 'worksheet': 'Calendar events'},
            {'model': 'calendarhub.CalendarNotificationPreference', 'worksheet': 'Calendar preferences'},
            {'model': 'calendarhub.CalendarReminderLog', 'worksheet': 'Calendar reminders'},
            {'model': 'calendarhub.CalendarAuditLog', 'worksheet': 'Calendar audits'},
        ],
    },
    {
        'module': 'vulnerabilities',
        'label': 'Vulnerability scan targets, schedules, jobs and findings',
        'formats': ['xlsx', 'csv_zip'],
        'sheets': [
            {'model': 'vuln.ScanTarget', 'worksheet': 'Scan targets'},
            {'model': 'vuln.ScanSchedule', 'worksheet': 'Scan schedules'},
            {'model': 'vuln.ScanJob', 'worksheet': 'Scan jobs'},
            {'model': 'vuln.VulnerabilityFinding', 'worksheet': 'Vulnerability findings'},
        ],
    },
]


def get_export_coverage_manifest():
    """Return documented customer data export coverage."""
    return EXPORT_COVERAGE


def _normalise_selected_modules(selected_modules):
    if not selected_modules or selected_modules == ['all']:
        return {entry['module'] for entry in EXPORT_COVERAGE}
    return set(selected_modules)


def _get_model(model_label):
    app_label, model_name = model_label.split('.', 1)
    return apps.get_model(app_label, model_name)


def _safe_sheet_title(title, used_titles):
    clean_title = title[:31]
    candidate = clean_title
    counter = 2
    while candidate in used_titles:
        suffix = f' {counter}'
        candidate = f'{clean_title[:31 - len(suffix)]}{suffix}'
        counter += 1
    used_titles.add(candidate)
    return candidate


def _exportable_fields(model_class):
    fields = []
    for field in model_class._meta.get_fields():
        if field.name in EXCLUDED_FIELD_NAMES:
            continue
        if field.auto_created and not field.concrete:
            continue
        if isinstance(field, models.ManyToManyField):
            fields.append(field)
            continue
        if not getattr(field, 'concrete', False):
            continue
        fields.append(field)
    return fields


def _serialise_value(value):
    if value is None:
        return ''
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, date | time):
        return value.isoformat()
    if isinstance(value, Decimal):
        return str(value)
    if isinstance(value, UUID):
        return str(value)
    if isinstance(value, dict | list | tuple):
        return json.dumps(value, default=str, ensure_ascii=False)
    return str(value)


def _row_for_instance(instance, fields):
    row = []
    for field in fields:
        if isinstance(field, models.ManyToManyField):
            manager = getattr(instance, field.name)
            value = ', '.join(str(item.pk) for item in manager.all())
        elif isinstance(field, models.ForeignKey):
            value = getattr(instance, f'{field.name}_id')
        elif isinstance(field, models.FileField):
            file_value = getattr(instance, field.name)
            value = file_value.name if file_value else ''
        else:
            value = getattr(instance, field.name)
        row.append(_serialise_value(value))
    return row


class TenantDataExportGenerator:
    """Generate tenant-scoped spreadsheet exports for all GRC data modules."""

    def __init__(self, data_export: TenantDataExport):
        self.data_export = data_export
        self.coverage = [
            entry for entry in EXPORT_COVERAGE
            if entry['module'] in _normalise_selected_modules(data_export.selected_modules)
        ]

    def generate_export(self):
        self.data_export.status = 'processing'
        self.data_export.generation_started_at = timezone.now()
        self.data_export.error_message = ''
        self.data_export.coverage_manifest = self.coverage
        self.data_export.save(update_fields=[
            'status',
            'generation_started_at',
            'error_message',
            'coverage_manifest',
        ])

        try:
            if self.data_export.export_format == 'xlsx':
                content, filename, mime_type, record_counts = self._generate_xlsx()
            elif self.data_export.export_format == 'csv_zip':
                content, filename, mime_type, record_counts = self._generate_csv_zip()
            else:
                raise ValueError(f'Unsupported export format: {self.data_export.export_format}')

            document = self._save_export_document(content, filename, mime_type)
            self.data_export.generated_file = document
            self.data_export.record_counts = record_counts
            self.data_export.status = 'completed'
            self.data_export.generation_completed_at = timezone.now()
            self.data_export.save(update_fields=[
                'generated_file',
                'record_counts',
                'status',
                'generation_completed_at',
            ])
            return document
        except Exception as exc:
            self.data_export.status = 'failed'
            self.data_export.error_message = str(exc)
            self.data_export.generation_completed_at = timezone.now()
            self.data_export.save(update_fields=[
                'status',
                'error_message',
                'generation_completed_at',
            ])
            raise

    def _iter_sheet_data(self):
        used_titles = set()
        for module in self.coverage:
            for sheet in module['sheets']:
                model_class = _get_model(sheet['model'])
                fields = _exportable_fields(model_class)
                headers = [field.name for field in fields]
                queryset = model_class.objects.all().order_by('pk')
                rows = [_row_for_instance(instance, fields) for instance in queryset]
                title = _safe_sheet_title(sheet['worksheet'], used_titles)
                yield title, headers, rows, sheet['model']

    def _generate_xlsx(self):
        workbook = Workbook()
        workbook.remove(workbook.active)
        record_counts = {}

        for title, headers, rows, model_label in self._iter_sheet_data():
            worksheet = workbook.create_sheet(title=title)
            worksheet.append(headers or ['id'])
            for row in rows:
                worksheet.append(row)
            record_counts[model_label] = len(rows)

        buffer = io.BytesIO()
        workbook.save(buffer)
        buffer.seek(0)
        timestamp = timezone.now().strftime('%Y%m%d_%H%M%S')
        return (
            buffer.getvalue(),
            f'tenant-data-export-{timestamp}.xlsx',
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            record_counts,
        )

    def _generate_csv_zip(self):
        zip_buffer = io.BytesIO()
        record_counts = {}
        with zipfile.ZipFile(zip_buffer, mode='w', compression=zipfile.ZIP_DEFLATED) as archive:
            for title, headers, rows, model_label in self._iter_sheet_data():
                csv_buffer = io.StringIO()
                writer = csv.writer(csv_buffer)
                writer.writerow(headers or ['id'])
                writer.writerows(rows)
                archive.writestr(f'{title}.csv', csv_buffer.getvalue())
                record_counts[model_label] = len(rows)

        zip_buffer.seek(0)
        timestamp = timezone.now().strftime('%Y%m%d_%H%M%S')
        return (
            zip_buffer.getvalue(),
            f'tenant-data-export-{timestamp}.zip',
            'application/zip',
            record_counts,
        )

    def _save_export_document(self, content, filename, mime_type):
        document = Document(
            title=self.data_export.title,
            description='Generated tenant data export',
            uploaded_by=self.data_export.requested_by,
            mime_type=mime_type,
        )
        document.file.save(filename, ContentFile(content), save=False)
        document.file_size = len(content)
        document.save()
        return document


class AssessmentReportGenerator:
    """Service for generating assessment reports in PDF format."""
    
    def __init__(self, report: AssessmentReport):
        self.report = report
        self.font_config = FontConfiguration()
    
    def generate_report(self):
        """Generate the PDF report based on report type."""
        try:
            self.report.status = 'processing'
            self.report.generation_started_at = timezone.now()
            self.report.save()
            
            # Generate report based on type
            if self.report.report_type == 'assessment_summary':
                html_content = self._generate_assessment_summary()
            elif self.report.report_type == 'detailed_assessment':
                html_content = self._generate_detailed_assessment()
            elif self.report.report_type == 'evidence_portfolio':
                html_content = self._generate_evidence_portfolio()
            elif self.report.report_type == 'compliance_gap':
                html_content = self._generate_compliance_gap()
            elif self.report.report_type == 'risk_analytics':
                html_content = self._generate_risk_analytics_report()
            else:
                raise ValueError(f"Unknown report type: {self.report.report_type}")
            
            # Generate PDF
            pdf_content = self._render_pdf(html_content)
            
            # Save PDF as Document
            filename = self._generate_filename()
            document = self._save_pdf_document(pdf_content, filename)
            
            # Update report status
            self.report.generated_file = document
            self.report.status = 'completed'
            self.report.generation_completed_at = timezone.now()
            self.report.save()
            
            return document
            
        except Exception as e:
            self.report.status = 'failed'
            self.report.error_message = str(e)
            self.report.generation_completed_at = timezone.now()
            self.report.save()
            raise
    
    def _generate_assessment_summary(self):
        """Generate assessment summary report HTML."""
        context = {
            'report': self.report,
            'framework': self.report.framework,
            'generated_at': timezone.now(),
        }
        
        if self.report.framework:
            # Framework-specific summary
            assessments = ControlAssessment.objects.filter(
                control__clauses__framework=self.report.framework
            ).select_related(
                'control', 'assigned_to'
            ).prefetch_related('control__clauses').distinct()
            
            # Assessment statistics
            total_assessments = assessments.count()
            completed_assessments = assessments.filter(status='complete').count()
            in_progress_assessments = assessments.filter(status='in_progress').count()
            not_started_assessments = assessments.filter(status='not_started').count()
            overdue_assessments = assessments.filter(
                status__in=['not_started', 'in_progress'],
                due_date__lt=timezone.now().date()
            ).count()
            
            # Evidence statistics
            evidence_stats = AssessmentEvidence.objects.filter(
                assessment__in=assessments
            ).aggregate(
                total_evidence=Count('id'),
                primary_evidence=Count('id', filter=Q(is_primary_evidence=True))
            )
            
            context.update({
                'assessments': assessments,
                'total_assessments': total_assessments,
                'completed_assessments': completed_assessments,
                'in_progress_assessments': in_progress_assessments,
                'not_started_assessments': not_started_assessments,
                'overdue_assessments': overdue_assessments,
                'completion_percentage': round((completed_assessments / total_assessments * 100) if total_assessments > 0 else 0, 1),
                'evidence_stats': evidence_stats,
            })
            
            # Add risk analytics data to provide comprehensive risk context
            context.update(self._get_risk_analytics_context())
        else:
            # All assessments summary
            assessments = self.report.assessments.all()
            context['assessments'] = assessments
        
        return render_to_string('exports/reports/assessment_summary.html', context)
    
    def _generate_detailed_assessment(self):
        """Generate detailed assessment report HTML."""
        context = {
            'report': self.report,
            'generated_at': timezone.now(),
        }
        
        if self.report.framework:
            assessments = ControlAssessment.objects.filter(
                control__clauses__framework=self.report.framework
            ).select_related(
                'control', 'assigned_to'
            ).prefetch_related('control__clauses', 'evidence_links__evidence').distinct()
        else:
            assessments = self.report.assessments.all().select_related(
                'control', 'assigned_to'
            ).prefetch_related('control__clauses', 'evidence_links__evidence')
        
        context['assessments'] = assessments
        return render_to_string('exports/reports/detailed_assessment.html', context)
    
    def _generate_evidence_portfolio(self):
        """Generate evidence portfolio report HTML."""
        context = {
            'report': self.report,
            'generated_at': timezone.now(),
        }
        
        # Get evidence across assessments
        if self.report.framework:
            evidence_links = AssessmentEvidence.objects.filter(
                assessment__control__clauses__framework=self.report.framework
            ).select_related(
                'assessment', 'evidence', 'assessment__control'
            ).prefetch_related('evidence__document').distinct()
        else:
            assessment_ids = self.report.assessments.values_list('id', flat=True)
            evidence_links = AssessmentEvidence.objects.filter(
                assessment_id__in=assessment_ids
            ).select_related(
                'assessment', 'evidence', 'assessment__control'
            ).prefetch_related('evidence__document')
        
        # Group evidence by type and reuse
        evidence_summary = {}
        for link in evidence_links:
            evidence = link.evidence
            if evidence.id not in evidence_summary:
                evidence_summary[evidence.id] = {
                    'evidence': evidence,
                    'assessments': [],
                    'is_primary_count': 0,
                }
            evidence_summary[evidence.id]['assessments'].append(link.assessment)
            if link.is_primary_evidence:
                evidence_summary[evidence.id]['is_primary_count'] += 1
        
        context['evidence_summary'] = evidence_summary.values()
        return render_to_string('exports/reports/evidence_portfolio.html', context)
    
    def _generate_compliance_gap(self):
        """Generate compliance gap analysis report HTML."""
        context = {
            'report': self.report,
            'framework': self.report.framework,
            'generated_at': timezone.now(),
        }
        
        if not self.report.framework:
            raise ValueError("Compliance gap analysis requires a framework")
        
        # Get all assessments for the framework
        assessments = ControlAssessment.objects.filter(
            control__clauses__framework=self.report.framework
        ).select_related(
            'control', 'assigned_to'
        ).prefetch_related('control__clauses').distinct()
        
        # Categorize gaps
        not_started = assessments.filter(status='not_started')
        in_progress_overdue = assessments.filter(
            status='in_progress',
            due_date__lt=timezone.now().date()
        )
        missing_evidence = assessments.filter(
            status='complete',
            evidence_links__isnull=True
        )
        no_primary_evidence = assessments.filter(
            status='complete'
        ).exclude(
            evidence_links__is_primary_evidence=True
        )
        
        context.update({
            'not_started': not_started,
            'in_progress_overdue': in_progress_overdue,
            'missing_evidence': missing_evidence,
            'no_primary_evidence': no_primary_evidence,
            'total_gaps': (
                not_started.count() + 
                in_progress_overdue.count() + 
                missing_evidence.count() + 
                no_primary_evidence.count()
            ),
        })
        
        return render_to_string('exports/reports/compliance_gap.html', context)
    
    def _render_pdf(self, html_content):
        """Convert HTML to PDF using WeasyPrint."""
        # Define CSS for styling
        css_content = """
        @page {
            size: A4;
            margin: 2cm;
            @bottom-center {
                content: "Page " counter(page) " of " counter(pages);
                font-size: 10pt;
                color: #666;
            }
        }
        
        body {
            font-family: 'Helvetica', 'Arial', sans-serif;
            font-size: 11pt;
            line-height: 1.4;
            color: #333;
        }
        
        .header {
            border-bottom: 2px solid #0066cc;
            padding-bottom: 1em;
            margin-bottom: 2em;
        }
        
        .header h1 {
            color: #0066cc;
            margin: 0;
            font-size: 20pt;
        }
        
        .header .subtitle {
            color: #666;
            font-size: 12pt;
            margin-top: 0.5em;
        }
        
        .stats-grid {
            display: grid;
            grid-template-columns: repeat(auto-fit, minmax(200px, 1fr));
            gap: 1em;
            margin: 1em 0;
        }
        
        .stat-card {
            border: 1px solid #ddd;
            border-radius: 4px;
            padding: 1em;
            text-align: center;
        }
        
        .stat-card .number {
            font-size: 24pt;
            font-weight: bold;
            color: #0066cc;
        }
        
        .stat-card .label {
            font-size: 10pt;
            color: #666;
            margin-top: 0.5em;
        }
        
        .assessment-table {
            width: 100%;
            border-collapse: collapse;
            margin: 1em 0;
        }
        
        .assessment-table th,
        .assessment-table td {
            border: 1px solid #ddd;
            padding: 8px;
            text-align: left;
            font-size: 10pt;
        }
        
        .assessment-table th {
            background-color: #f5f5f5;
            font-weight: bold;
        }
        
        .status-badge {
            padding: 2px 8px;
            border-radius: 4px;
            font-size: 9pt;
            font-weight: bold;
            color: white;
        }
        
        .status-complete,
        .status-completed { background-color: #28a745; }
        .status-in-progress { background-color: #ffc107; color: #333; }
        .status-not-started { background-color: #6c757d; }
        .status-overdue { background-color: #dc3545; }
        
        .page-break {
            page-break-before: always;
        }
        
        .section {
            margin: 2em 0;
        }
        
        .section h2 {
            color: #0066cc;
            border-bottom: 1px solid #0066cc;
            padding-bottom: 0.5em;
        }
        """
        
        css = CSS(string=css_content, font_config=self.font_config)
        html = HTML(string=html_content)
        
        pdf_buffer = io.BytesIO()
        html.write_pdf(
            pdf_buffer,
            stylesheets=[css],
            font_config=self.font_config,
            presentational_hints=False,
        )
        pdf_buffer.seek(0)
        
        return pdf_buffer.getvalue()
    
    def _generate_filename(self):
        """Generate appropriate filename for the report."""
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        report_type = self.report.report_type.replace('_', '-')
        
        if self.report.framework:
            framework_name = self.report.framework.short_name.lower().replace(' ', '-')
            return f"{report_type}-{framework_name}-{timestamp}.pdf"
        else:
            return f"{report_type}-{timestamp}.pdf"
    
    def _save_pdf_document(self, pdf_content, filename):
        """Save PDF content as a Document."""
        document = Document(
            title=f"{self.report.title}",
            description=f"Generated assessment report: {self.report.get_report_type_display()}",
            document_type='report',
            file_category='generated',
            uploaded_by=self.report.requested_by
        )
        
        # Save the PDF content
        document.file.save(
            filename,
            ContentFile(pdf_content),
            save=False
        )
        
        # Set file metadata
        document.file_size = len(pdf_content)
        document.mime_type = 'application/pdf'
        document.save()
        
        return document
    
    def _get_risk_analytics_context(self):
        """
        Get risk analytics data to include in assessment reports.
        Provides comprehensive risk context to complement compliance data.
        """
        try:
            # Get risk overview data
            risk_overview = RiskAnalyticsService.get_risk_overview_stats()
            
            # Get risk action progress data
            action_overview = RiskAnalyticsService.get_risk_action_overview_stats()
            
            # Get executive risk summary for high-level insights
            executive_summary = RiskAnalyticsService.get_executive_risk_summary()
            
            # Get heat map data for risk visualization context
            heat_map_data = RiskAnalyticsService.get_risk_heat_map_data()
            
            # Calculate risk-compliance correlation metrics
            risk_compliance_metrics = self._calculate_risk_compliance_correlation()
            
            return {
                'risk_analytics': {
                    'overview': risk_overview,
                    'actions': action_overview,
                    'executive_summary': executive_summary,
                    'heat_map': heat_map_data,
                    'compliance_correlation': risk_compliance_metrics,
                }
            }
        except Exception as e:
            # Graceful degradation - if risk analytics fail, don't break the report
            return {
                'risk_analytics': {
                    'error': f"Risk analytics unavailable: {str(e)}",
                    'overview': {'total_risks': 0, 'active_risks': 0},
                    'actions': {'total_actions': 0, 'completed_actions': 0},
                }
            }
    
    def _calculate_risk_compliance_correlation(self):
        """
        Calculate correlation metrics between risk management and compliance activities.
        This provides insights into how risk mitigation aligns with compliance efforts.
        """
        try:
            # Count risks that have associated compliance actions
            total_risks = Risk.objects.count()
            risks_with_actions = Risk.objects.filter(riskaction__isnull=False).distinct().count()
            
            # Count overdue risks vs overdue assessments for comparative analysis
            overdue_risks = Risk.objects.filter(
                review_date__lt=timezone.now().date(),
                status__in=['identified', 'assessed', 'treatment_planned']
            ).count()
            
            # Calculate action completion rate for risk mitigation
            total_risk_actions = RiskAction.objects.count()
            completed_risk_actions = RiskAction.objects.filter(status='completed').count()
            
            risk_action_completion_rate = (
                completed_risk_actions / total_risk_actions * 100
            ) if total_risk_actions > 0 else 0
            
            return {
                'total_risks': total_risks,
                'risks_with_mitigation_actions': risks_with_actions,
                'risk_mitigation_coverage': (
                    risks_with_actions / total_risks * 100
                ) if total_risks > 0 else 0,
                'overdue_risks': overdue_risks,
                'total_risk_actions': total_risk_actions,
                'completed_risk_actions': completed_risk_actions,
                'risk_action_completion_rate': round(risk_action_completion_rate, 1),
                'integration_insights': {
                    'risk_driven_compliance': risks_with_actions > 0,
                    'action_based_mitigation': total_risk_actions > 0,
                    'comprehensive_coverage': (
                        risks_with_actions / total_risks > 0.7
                    ) if total_risks > 0 else False,
                }
            }
        except Exception:
            return {
                'total_risks': 0,
                'risks_with_mitigation_actions': 0,
                'risk_mitigation_coverage': 0,
                'integration_insights': {
                    'risk_driven_compliance': False,
                    'action_based_mitigation': False,
                    'comprehensive_coverage': False,
                }
            }
    
    def _generate_risk_analytics_report(self):
        """Generate comprehensive risk analytics report HTML."""
        context = {
            'report': self.report,
            'generated_at': timezone.now(),
        }
        
        # Get comprehensive risk analytics data
        risk_context = self._get_risk_analytics_context()
        context.update(risk_context)
        
        # Get additional risk analytics data for detailed report
        try:
            context.update({
                'trend_analysis': RiskAnalyticsService.get_risk_trend_analysis(),
                'control_integration': RiskAnalyticsService.get_risk_control_integration_analysis(),
                'dashboard_data': RiskReportGenerator.generate_risk_dashboard_data(),
            })
        except Exception as e:
            context['analytics_error'] = str(e)
        
        return render_to_string('exports/reports/risk_analytics.html', context)
